"""The plan sheet: a real .xlsx that the planning agent writes and the deploy agent reads.

One row per finalized decision. This is the hand-off contract between the two agents,
and the version-controlled source of truth for the rollout. Kept as a spreadsheet so a
human can open, review, and share it.
"""
from __future__ import annotations
import os
import json
import datetime as _dt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

PLAN_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "plan.xlsx")

COLUMNS = [
    "id", "area", "title", "value_json", "pros", "watchouts", "one_way",
    "rationale", "source_url", "terraform_resource", "owner", "status",
    "terraform_path", "deploy_status", "updated_at",
]

_HEADER_FILL = PatternFill("solid", fgColor="1967D2")
_ONEWAY_FILL = PatternFill("solid", fgColor="FCE8E6")


def _ensure(path: str = PLAN_PATH):
    if os.path.exists(path):
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "plan"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HEADER_FILL
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    widths = [8, 8, 26, 40, 34, 40, 8, 26, 34, 34, 16, 16, 30, 18, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    wb.save(path)


def save_decision(decision: dict, owner: str = "planning-agent", path: str = PLAN_PATH) -> dict:
    """Append a finalized decision to the plan sheet.

    Args:
      decision: a recommendation dict (from recommend.py) the user approved.
      owner: who owns this decision.

    Returns:
      {"id", "row", "status"} for the saved row.
    """
    _ensure(path)
    wb = load_workbook(path)
    ws = wb["plan"]
    next_id = ws.max_row  # header is row 1, so first data row id = 1
    row = [
        f"D-{next_id:03d}",
        decision.get("area", ""),
        decision.get("title", ""),
        json.dumps(decision.get("value", {}), ensure_ascii=False),
        " | ".join(decision.get("pros", [])),
        " | ".join(decision.get("watchouts", [])),
        "YES" if decision.get("one_way") else "no",
        decision.get("rationale", ""),
        decision.get("source_url", ""),
        decision.get("terraform_resource", ""),
        owner,
        "saved",
        "",  # terraform_path
        "",  # deploy_status
        _dt.datetime.utcnow().isoformat(timespec="seconds") + "Z",
    ]
    ws.append(row)
    if decision.get("one_way"):
        for c in ws[ws.max_row]:
            c.fill = _ONEWAY_FILL
    wb.save(path)
    return {"id": row[0], "row": ws.max_row, "status": "saved", "plan_path": path}


def read_decisions(path: str = PLAN_PATH, status: str | None = None) -> list[dict]:
    """Read decisions from the plan sheet (optionally filtered by status)."""
    if not os.path.exists(path):
        return []
    wb = load_workbook(path, read_only=True)
    ws = wb["plan"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    out = []
    for r in rows:
        rec = dict(zip(COLUMNS, r))
        if rec.get("id") is None:
            continue
        try:
            rec["value"] = json.loads(rec.get("value_json") or "{}")
        except Exception:
            rec["value"] = {}
        if status is None or rec.get("status") == status:
            out.append(rec)
    return out


def update_status(decision_id: str, status: str | None = None,
                  deploy_status: str | None = None, terraform_path: str | None = None,
                  path: str = PLAN_PATH) -> bool:
    """Update status / deploy_status / terraform_path for a decision (used by deploy agent)."""
    if not os.path.exists(path):
        return False
    wb = load_workbook(path)
    ws = wb["plan"]
    idx = {name: i + 1 for i, name in enumerate(COLUMNS)}
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=idx["id"]).value == decision_id:
            if status is not None:
                ws.cell(row=r, column=idx["status"]).value = status
            if deploy_status is not None:
                ws.cell(row=r, column=idx["deploy_status"]).value = deploy_status
            if terraform_path is not None:
                ws.cell(row=r, column=idx["terraform_path"]).value = terraform_path
            ws.cell(row=r, column=idx["updated_at"]).value = \
                _dt.datetime.utcnow().isoformat(timespec="seconds") + "Z"
            wb.save(path)
            return True
    return False
